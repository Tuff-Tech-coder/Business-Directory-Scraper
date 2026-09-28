"""Regression coverage for data loss, malformed input, and export safety."""

import importlib
from urllib.parse import parse_qs, urlsplit

import openpyxl
import pytest
from bs4 import BeautifulSoup

import directory_scraper as ds


def listing(extra="", name="Example"):
    return BeautifulSoup(
        f'<div class="result"><a class="business-name">{name}</a>{extra}</div>', "lxml"
    ).div


def test_preserves_functional_website_query():
    biz = ds.parse_listing(
        listing(
            '<a class="track-visit-website" href="https://example.com/?location=2&amp;utm_source=yp">Visit</a>'
        )
    )
    assert parse_qs(urlsplit(biz.website).query) == {"location": ["2"]}


def test_half_star_rating():
    assert (
        ds.parse_listing(listing('<div class="result-rating four half"></div>')).rating == "4.5/5"
    )


def test_text_nodes_keep_spaces():
    assert (
        ds.parse_listing(listing(name="<span>Ace</span> <span>Plumbing</span>")).name
        == "Ace Plumbing"
    )


def test_query_relative_pagination():
    soup = BeautifulSoup('<a class="next" href="?page=2">Next</a>', "lxml")
    assert ds.parse_results_page(soup)[1] == "https://www.yellowpages.com/search?page=2"


def test_protocol_relative_pagination():
    soup = BeautifulSoup(
        '<a class="next" href="//www.yellowpages.com/search?page=2">Next</a>', "lxml"
    )
    assert ds.parse_results_page(soup)[1] == "https://www.yellowpages.com/search?page=2"


def test_invalid_page_limit():
    with pytest.raises(ValueError):
        ds.DirectoryScraper("plumbers", "Austin, TX").scrape(0)


def test_nested_output_directory(tmp_path):
    path = tmp_path / "exports" / "results.xlsx"
    ds.export_to_excel([ds.Business(name="Example")], path)
    assert path.is_file()


def test_excel_formula_is_literal_text(tmp_path):
    path = tmp_path / "results.xlsx"
    ds.export_to_excel([ds.Business(name='=HYPERLINK("https://example.com","Click")')], path)
    ws = openpyxl.load_workbook(path)["Results"]
    assert ws["A2"].data_type == "s"
    assert ws["A2"].value.startswith("=HYPERLINK")


def test_excel_control_characters_do_not_abort(tmp_path):
    path = tmp_path / "results.xlsx"
    ds.export_to_excel([ds.Business(name="Ace\x01 Plumbing")], path)
    assert openpyxl.load_workbook(path)["Results"]["A2"].value == "Ace Plumbing"


def test_import_does_not_create_log(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    importlib.reload(ds)
    assert not (tmp_path / "directory_scraper.log").exists()


def test_demo_websites_use_reserved_domains():
    results = ds.generate_demo_results("plumbers", "Austin, TX", 100)
    assert all(urlsplit(b.website).hostname.endswith(".example") for b in results if b.website)


def test_demo_city_does_not_use_texas_zip_codes():
    results = ds.generate_demo_results("dentists", "Chicago, IL", 10)
    assert all("787" not in b.city_state for b in results)
