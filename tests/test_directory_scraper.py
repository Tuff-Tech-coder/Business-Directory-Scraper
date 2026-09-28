"""Tests for parsing and export.

Every test runs against saved HTML fixtures -- no network access, so the suite
is deterministic and runs in CI. The live-fetch path is deliberately not
exercised here; parsing and export are the parts that carry the logic.
"""

import sys
from pathlib import Path

import pytest
from bs4 import BeautifulSoup

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from fixtures import (  # noqa: E402
    AD_CARD,
    FULL_CARD,
    INTERNAL_REDIRECT_CARD,
    MINIMAL_CARD,
    NAMELESS_CARD,
    page,
)

from directory_scraper import (  # noqa: E402
    Business,
    export_to_excel,
    parse_listing,
    parse_results_page,
)


def card(html: str) -> BeautifulSoup:
    """Parse fixture markup and return the listing card element."""
    return BeautifulSoup(html, "lxml").find("div", class_="result")


def soup(html: str) -> BeautifulSoup:
    return BeautifulSoup(html, "lxml")


# ---------------------------------------------------------------------------
# parse_listing -- defensive extraction
# ---------------------------------------------------------------------------
class TestParseListing:
    def test_extracts_every_field(self):
        biz = parse_listing(card(FULL_CARD))
        assert biz.name == "Ace Plumbing & Drain"
        assert biz.phone == "(512) 555-0142"
        assert biz.address == "1420 Commerce Blvd"
        assert biz.city_state == "Austin, TX 78701"
        assert biz.category == "Plumbers, Drain Cleaning"

    def test_strips_tracking_params_from_website(self):
        biz = parse_listing(card(FULL_CARD))
        assert biz.website == "https://aceplumbing.com/"
        assert "utm_source" not in biz.website

    def test_rating_class_word_maps_to_number(self):
        biz = parse_listing(card(FULL_CARD))
        assert biz.rating == "4/5"

    def test_review_count_strips_non_digits(self):
        biz = parse_listing(card(FULL_CARD))
        assert biz.review_count == "87"

    def test_missing_elements_yield_empty_strings_not_errors(self):
        """The defensive-extraction case: absent markup must not raise."""
        biz = parse_listing(card(MINIMAL_CARD))
        assert biz is not None
        assert biz.name == "Bare Minimum Plumbing"
        assert biz.phone == ""
        assert biz.address == ""
        assert biz.city_state == ""
        assert biz.website == ""
        assert biz.category == ""
        assert biz.rating == ""
        assert biz.review_count == ""

    def test_returns_none_without_a_business_name(self):
        assert parse_listing(card(NAMELESS_CARD)) is None

    def test_internal_redirect_is_not_treated_as_a_website(self):
        """Only absolute external URLs count; /redirect?to=... is not a site."""
        biz = parse_listing(card(INTERNAL_REDIRECT_CARD))
        assert biz.website == ""

    def test_date_scraped_is_populated_automatically(self):
        biz = parse_listing(card(FULL_CARD))
        assert biz.date_scraped  # __post_init__ fills today's date

    def test_unrated_listing_has_empty_rating_not_zero(self):
        """An absent rating must not render as '0/5', which would be a claim."""
        biz = parse_listing(card(MINIMAL_CARD))
        assert biz.rating == ""


# ---------------------------------------------------------------------------
# parse_results_page -- ad filtering and pagination
# ---------------------------------------------------------------------------
class TestParseResultsPage:
    def test_filters_out_paid_placements(self):
        listings, _ = parse_results_page(soup(page(FULL_CARD + AD_CARD)))
        names = [b.name for b in listings]
        assert "Ace Plumbing & Drain" in names
        assert "Sponsored Plumbing Co" not in names
        assert len(listings) == 1

    def test_keeps_every_organic_listing(self):
        listings, _ = parse_results_page(soup(page(FULL_CARD + MINIMAL_CARD)))
        assert len(listings) == 2

    def test_drops_cards_that_fail_to_parse(self):
        listings, _ = parse_results_page(soup(page(FULL_CARD + NAMELESS_CARD)))
        assert len(listings) == 1

    def test_builds_absolute_next_url_from_relative_href(self):
        _, next_url = parse_results_page(soup(page(FULL_CARD, "/search?page=2")))
        assert next_url == "https://www.yellowpages.com/search?page=2"

    def test_absolute_next_url_is_left_alone(self):
        _, next_url = parse_results_page(soup(page(FULL_CARD, "https://example.com/search?page=3")))
        assert next_url == "https://example.com/search?page=3"

    def test_last_page_reports_no_next_url(self):
        _, next_url = parse_results_page(soup(page(FULL_CARD, next_href=None)))
        assert next_url is None

    def test_empty_page_returns_no_listings_and_no_next(self):
        listings, next_url = parse_results_page(soup(page("", next_href=None)))
        assert listings == []
        assert next_url is None

    def test_page_of_only_ads_yields_nothing(self):
        listings, _ = parse_results_page(soup(page(AD_CARD + AD_CARD)))
        assert listings == []


# ---------------------------------------------------------------------------
# export_to_excel
# ---------------------------------------------------------------------------
class TestExportToExcel:
    @pytest.fixture
    def businesses(self):
        return [
            Business(
                name="Ace Plumbing",
                phone="(512) 555-0142",
                address="1 Main St",
                city_state="Austin, TX 78701",
                website="https://ace.com",
                category="Plumbers",
                rating="4/5",
                review_count="87",
            ),
            Business(
                name="Bee Plumbing",
                phone="(512) 555-0199",
                address="2 Oak Ave",
                city_state="Austin, TX 78702",
            ),
        ]

    def test_writes_a_readable_workbook(self, businesses, tmp_path):
        import openpyxl

        out = tmp_path / "results.xlsx"
        export_to_excel(businesses, str(out))
        assert out.exists()

        ws = openpyxl.load_workbook(out)["Results"]
        assert ws.max_row == 3  # header + 2 rows
        assert ws.cell(row=2, column=1).value == "Ace Plumbing"
        assert ws.cell(row=3, column=1).value == "Bee Plumbing"

    def test_header_row_matches_expected_columns(self, businesses, tmp_path):
        import openpyxl

        out = tmp_path / "results.xlsx"
        export_to_excel(businesses, str(out))
        ws = openpyxl.load_workbook(out)["Results"]
        headers = [c.value for c in ws[1]]
        assert headers == [
            "Business Name",
            "Phone",
            "Address",
            "City / State / ZIP",
            "Website",
            "Category",
            "Rating",
            "Review Count",
            "Date Scraped",
        ]

    def test_empty_input_writes_no_file(self, tmp_path):
        """No results is a warning, not a crash, and not an empty workbook."""
        out = tmp_path / "empty.xlsx"
        export_to_excel([], str(out))
        assert not out.exists()

    def test_partial_records_export_as_blanks(self, businesses, tmp_path):
        import openpyxl

        out = tmp_path / "results.xlsx"
        export_to_excel(businesses, str(out))
        ws = openpyxl.load_workbook(out)["Results"]
        # Bee Plumbing has no website/category/rating -- those cells are blank.
        assert ws.cell(row=3, column=5).value in (None, "")
        assert ws.cell(row=3, column=7).value in (None, "")
