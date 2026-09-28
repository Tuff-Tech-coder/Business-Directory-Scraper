"""Exercise transport, pagination, CLI exits, and real dashboard HTTP responses."""

import io
import json
import threading
from http.client import HTTPConnection
from http.server import ThreadingHTTPServer
from types import SimpleNamespace
from unittest.mock import Mock

import openpyxl
import pytest
from bs4 import BeautifulSoup

import directory_scraper as ds
from dashboard import DemoHandler


def page(name="Example", next_href=None):
    html = f'<div class="result"><a class="business-name">{name}</a></div>'
    if next_href:
        html += f'<a class="next" href="{next_href}">Next</a>'
    return BeautifulSoup(html, "lxml")


def response(status, text):
    return SimpleNamespace(status_code=status, text=text, content=text.encode())


def test_small_valid_page_is_accepted():
    session = Mock()
    session.get.return_value = response(200, str(page()))
    assert ds.fetch_page("https://example.com", session) is not None
    assert session.get.call_count == 1


@pytest.mark.parametrize("status", [401, 403, 404])
def test_permanent_http_error_is_not_retried(status):
    session = Mock()
    session.get.return_value = response(status, "denied")
    assert ds.fetch_page("https://example.com", session) is None
    assert session.get.call_count == 1


def test_transient_failure_retries_with_backoff(monkeypatch):
    session, sleep = Mock(), Mock()
    session.get.side_effect = [response(503, "busy"), response(200, str(page()))]
    monkeypatch.setattr(ds.time, "sleep", sleep)
    assert ds.fetch_page("https://example.com", session) is not None
    sleep.assert_called_once_with(5)
    assert session.get.call_count == 2


def test_large_interstitial_does_not_become_empty_success():
    session = Mock()
    session.get.return_value = response(200, "<html>" + "Challenge " * 2000 + "</html>")
    assert ds.fetch_page("https://example.com", session) is None


def test_pagination_loop_stops_and_duplicates_are_removed():
    scraper = ds.DirectoryScraper("plumbers", "Austin, TX")
    second = scraper._first_url() + "&page=2"
    loader = Mock(side_effect=[page(next_href=second), page(next_href=second)])
    result = scraper.scrape(10, page_loader=loader, throttle=False)
    assert len(result) == 1
    assert loader.call_count == 2
    assert scraper.duplicates_skipped == 1
    assert "loop" in scraper.last_error


def test_cross_host_pagination_is_not_fetched():
    scraper = ds.DirectoryScraper("plumbers", "Austin, TX")
    loader = Mock(return_value=page(next_href="https://example.com/page2"))
    assert len(scraper.scrape(5, page_loader=loader, throttle=False)) == 1
    assert loader.call_count == 1
    assert "outside" in scraper.last_error


def test_failed_later_page_preserves_results_and_failure():
    scraper = ds.DirectoryScraper("plumbers", "Austin, TX")
    loader = Mock(side_effect=[page(next_href="?page=2"), None])
    assert len(scraper.scrape(5, page_loader=loader, throttle=False)) == 1
    assert "page 2" in scraper.last_error


def test_context_manager_closes_http_session():
    session = Mock()
    with ds.DirectoryScraper("plumbers", "Austin, TX") as scraper:
        scraper.session = session
    session.close.assert_called_once()


def test_demo_walks_parser_without_creating_http_session(monkeypatch):
    monkeypatch.setattr(
        ds.cffi_requests, "Session", Mock(side_effect=AssertionError("Network session created"))
    )
    businesses, stats = ds.run_demo()
    assert len(businesses) == 35
    assert stats["pages"] == 3
    assert stats["duplicates"] == 2
    assert stats["ads"] == 6
    assert sum(e["added"] for e in stats["events"]) == 35
    assert businesses == ds.generate_demo_results("plumbers", "Austin, TX")


def test_demo_does_not_change_global_random_state():
    before = ds.random.getstate()
    ds.run_demo()
    assert ds.random.getstate() == before


@pytest.mark.parametrize(
    "args",
    [["--max-pages", "0"], ["--category", " "], ["--count", "1001"], ["--output", "file.csv"]],
)
def test_cli_rejects_invalid_arguments(args):
    with pytest.raises(SystemExit) as exc:
        ds.main(args)
    assert exc.value.code == 2


def test_cli_returns_failure_for_empty_results(monkeypatch):
    monkeypatch.setattr(ds.DirectoryScraper, "scrape", lambda *a, **k: [])
    assert ds.main([]) == 1


def test_cli_partial_run_is_nonzero_and_marked_in_workbook(tmp_path, monkeypatch, capsys):
    def partial(self, max_pages):
        self.last_error = "Page 2 unavailable"
        return [ds.Business(name="Kept record")]

    monkeypatch.setattr(ds.DirectoryScraper, "scrape", partial)
    out = tmp_path / "partial.xlsx"
    assert ds.main(["--output", str(out)]) == 2
    assert "[PARTIAL]" in capsys.readouterr().out
    assert "PARTIAL" in openpyxl.load_workbook(out)["About this export"]["B2"].value


def test_excel_metrics_and_date_are_typed():
    buffer = io.BytesIO()
    ds.export_to_excel([ds.Business(name="Example", rating="4.5/5", review_count="120")], buffer)
    ws = openpyxl.load_workbook(buffer)["Results"]
    assert ws["G2"].value == 4.5
    assert ws["H2"].value == 120
    assert ws["I2"].is_date
    assert ws.freeze_panes == "B2"
    assert ws.auto_filter.ref == "A1:I2"


@pytest.fixture
def dashboard():
    server = ThreadingHTTPServer(("127.0.0.1", 0), DemoHandler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    yield server.server_port
    server.shutdown()
    server.server_close()
    thread.join()


def get(port, path, headers=None):
    conn = HTTPConnection("127.0.0.1", port, timeout=5)
    try:
        conn.request("GET", path, headers=headers or {})
        resp = conn.getresponse()
        return resp.status, resp.read(), dict(resp.getheaders())
    finally:
        conn.close()


def test_dashboard_serves_dataset_and_actual_excel(dashboard):
    status, body, _ = get(dashboard, "/api/demo")
    assert status == 200
    payload = json.loads(body)
    assert len(payload["businesses"]) == 35
    status, body, headers = get(dashboard, "/api/export?rating=4&website=true")
    assert status == 200
    assert "attachment" in headers["Content-Disposition"]
    ws = openpyxl.load_workbook(io.BytesIO(body))["Results"]
    expected = [
        b
        for b in payload["businesses"]
        if b["rating"] and float(b["rating"].split("/")[0]) >= 4 and b["website"]
    ]
    assert ws.max_row == len(expected) + 1
    assert [ws.cell(i + 2, 1).value for i in range(len(expected))] == [b["name"] for b in expected]


@pytest.mark.parametrize(
    "path",
    [
        "/api/demo?count=-1",
        "/api/demo?count=no",
        "/api/demo?category=bad",
        "/api/export?rating=nan",
        "/api/export?search=no-such-business",
    ],
)
def test_dashboard_invalid_queries_report_400(dashboard, path):
    assert get(dashboard, path)[0] == 400


def test_dashboard_rejects_foreign_host_and_arbitrary_paths(dashboard):
    assert get(dashboard, "/api/demo", {"Host": "evil.example"})[0] == 403
    assert get(dashboard, "/../directory_scraper.py")[0] == 404


def test_dashboard_static_files(dashboard):
    for path in ["/", "/app.css", "/app.js"]:
        status, body, headers = get(dashboard, path)
        assert status == 200
        assert body
        assert "Content-Security-Policy" in headers
