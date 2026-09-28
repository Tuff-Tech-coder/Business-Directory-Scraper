"""
Business Directory Scraper
==========================
A structured data-extraction pipeline that turns paginated directory listings
into a formatted Excel workbook. Extracts business name, phone, address,
website, rating and category, walks pagination automatically, and throttles
requests with randomised delays.

The parser targets Yellow Pages-style HTML and is tested against fixtures.
The Business dataclass and Excel exporter can be reused with other sources.

Responsible use
---------------
The live-fetch path targets Yellow Pages. Check a site's Terms of Service and
robots.txt before pointing this at it, and prefer an official API or a
licensed data provider where one exists. --demo runs the entire pipeline
against generated data with no network access, and is the intended path for
evaluating this project.

Usage:
    python directory_scraper.py --serve   # local offline dashboard
    python directory_scraper.py --demo    # full pipeline, no network requests
    python directory_scraper.py --category "plumbers" --city "Austin, TX"
    python directory_scraper.py --category "dentists" --city "Chicago, IL" --max-pages 5

Output: A formatted Excel file with all results.
"""

import argparse
import datetime
import io
import logging
import os
import random
import re
import tempfile
import time
from collections.abc import Callable
from contextlib import suppress
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import parse_qsl, urlencode, urljoin, urlsplit, urlunsplit

import pandas as pd
from bs4 import BeautifulSoup
from curl_cffi import requests as cffi_requests
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

# ---------------------------------------------------------------------------
# Logging
# ---------------------------------------------------------------------------
logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------
@dataclass
class Business:
    name: str = ""
    phone: str = ""
    address: str = ""
    city_state: str = ""
    website: str = ""
    category: str = ""
    rating: str = ""
    review_count: str = ""
    date_scraped: str = ""

    def __post_init__(self):
        if not self.date_scraped:
            self.date_scraped = datetime.date.today().isoformat()


# ---------------------------------------------------------------------------
# HTTP fetch
# ---------------------------------------------------------------------------
def fetch_page(url: str, session: cffi_requests.Session, retries: int = 3) -> BeautifulSoup | None:
    """
    Fetch a URL and return a BeautifulSoup object, or None on failure.

    Validate recognizable listing/empty-results markup rather than body size.
    Retry transient failures; stop immediately on access denials and other
    permanent HTTP errors. An interstitial is not a successful empty search.

    Returns None rather than raising: the caller decides whether a failed page
    ends the run, so one bad page cannot abort a long pagination walk.
    """
    if retries < 1:
        raise ValueError("retries must be positive")
    for attempt in range(1, retries + 1):
        try:
            resp = session.get(
                url,
                headers={
                    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
                    "Accept-Language": "en-US,en;q=0.9",
                    "Referer": "https://www.google.com/",
                    "Upgrade-Insecure-Requests": "1",
                },
                timeout=20,
            )
            if resp.status_code == 200:
                soup = BeautifulSoup(resp.text, "lxml")
                if soup.select_one(".result, .no-results, .no-results-message"):
                    return soup
                logger.warning("Response has no recognizable results markup")
                return None
            if 400 <= resp.status_code < 500 and resp.status_code not in (408, 429):
                logger.warning("Request refused: HTTP %s", resp.status_code)
                return None
            logger.warning(
                f"Unexpected response on attempt {attempt}: status={resp.status_code} size={len(resp.content)}"
            )
        except Exception as e:
            logger.warning(f"Request error on attempt {attempt}: {e}")

        if attempt < retries:
            wait = 5 * attempt
            logger.info(f"Waiting {wait}s before retry...")
            time.sleep(wait)

    return None


# ---------------------------------------------------------------------------
# Parser — uses selectors confirmed against live Yellow Pages HTML
# ---------------------------------------------------------------------------
def clean_website(href: str) -> str:
    """Keep functional URL parameters; discard only known tracking parameters."""
    try:
        parts = urlsplit(href.strip())
        if parts.scheme not in {"http", "https"} or not parts.hostname:
            return ""
        if parts.username or parts.password:
            return ""
        params = [
            (key, value)
            for key, value in parse_qsl(parts.query, keep_blank_values=True)
            if not key.lower().startswith("utm_") and key.lower() not in {"ref", "gclid", "fbclid"}
        ]
        return urlunsplit(parts._replace(query=urlencode(params)))
    except ValueError:
        return ""


def parse_listing(card: BeautifulSoup) -> Business | None:
    """Parse a single Yellow Pages organic listing card."""

    def safe_text(tag_name: str, attrs: dict) -> str:
        el = card.find(tag_name, attrs)
        return el.get_text(" ", strip=True) if el else ""

    name = safe_text("a", {"class": "business-name"})
    if not name:
        return None

    # Phone is in a div with classes "phones phone primary"
    phone_el = card.find("div", class_="phones")
    phone = phone_el.get_text(strip=True) if phone_el else ""

    # Address: street is div.street-address, city/state is div.locality
    street = safe_text("div", {"class": "street-address"})
    locality = safe_text("div", {"class": "locality"})

    # Website — the actual external URL is in the href attribute
    website = ""
    ws_tag = card.find("a", class_="track-visit-website")
    if ws_tag:
        website = clean_website(ws_tag.get("href", ""))

    # Categories
    cats_el = card.find("div", class_="categories")
    category = cats_el.get_text(separator=", ", strip=True) if cats_el else ""

    # Rating — star level is encoded as a CSS class like "result-rating three"
    rating = ""
    rating_el = card.find("div", class_="result-rating")
    if rating_el:
        classes = rating_el.get("class", [])
        word_map = {"one": "1", "two": "2", "three": "3", "four": "4", "five": "5"}
        for cls in classes:
            if cls in word_map:
                rating = word_map[cls]
                if "half" in classes and rating != "5":
                    rating = f"{int(rating) + 0.5:g}"
                break

    # Review count is in span.count inside the ratings div
    count_el = card.find("span", class_="count")
    review_count = re.sub(r"[^\d]", "", count_el.get_text()) if count_el else ""

    return Business(
        name=name,
        phone=phone,
        address=street,
        city_state=locality,
        website=website,
        category=category,
        rating=f"{rating}/5" if rating else "",
        review_count=review_count,
    )


def parse_results_page(
    soup: BeautifulSoup, base_url: str = "https://www.yellowpages.com/search"
) -> tuple[list[Business], str | None]:
    """
    Parse all organic listings from a page.
    Returns (listings, next_page_url).
    """
    # Organic results don't have the 'flash-endt' ad class
    all_cards = soup.find_all("div", class_="result")
    organic_cards = [c for c in all_cards if "flash-endt" not in c.get("class", [])]

    listings = []
    for card in organic_cards:
        biz = parse_listing(card)
        if biz and biz.name:
            listings.append(biz)

    # Next page link
    next_link = soup.find("a", class_="next")
    next_url = None
    if next_link and next_link.get("href"):
        href = next_link["href"]
        candidate = urljoin(base_url, href)
        try:
            parts = urlsplit(candidate)
            if parts.scheme in {"http", "https"} and parts.hostname:
                next_url = urlunsplit(parts._replace(fragment=""))
        except ValueError:
            pass

    return listings, next_url


# ---------------------------------------------------------------------------
# Main scraper
# ---------------------------------------------------------------------------
class DirectoryScraper:
    SEARCH_URL = "https://www.yellowpages.com/search"

    def __init__(self, category: str, city: str):
        self.category = category.strip()
        self.city = city.strip()
        if not self.category or not self.city:
            raise ValueError("Category and city must not be blank")
        self.session = None
        self.last_error = None
        self.pages_scraped = 0
        self.duplicates_skipped = 0
        self.events = []

    def close(self):
        if self.session is not None:
            self.session.close()
            self.session = None

    def __enter__(self):
        return self

    def __exit__(self, *_):
        self.close()

    def _first_url(self) -> str:
        params = {"search_terms": self.category, "geo_location_terms": self.city}
        return f"{self.SEARCH_URL}?{urlencode(params)}"

    def scrape(
        self,
        max_pages: int = 10,
        *,
        page_loader: Callable[[str], BeautifulSoup | None] | None = None,
        throttle: bool = True,
    ) -> list[Business]:
        if max_pages < 1:
            raise ValueError("max_pages must be positive")
        self.last_error = None
        self.pages_scraped = self.duplicates_skipped = 0
        self.events = []
        all_results: list[Business] = []
        url = self._first_url()
        visited = set()
        seen = set()

        def load(url):
            if self.session is None:
                self.session = cffi_requests.Session(impersonate="chrome124")
            return fetch_page(url, self.session)

        loader = page_loader or load

        for page_num in range(1, max_pages + 1):
            logger.info(f"Scraping page {page_num}: {url}")

            visited.add(url)
            soup = loader(url)
            if soup is None:
                self.last_error = f"Could not load page {page_num}; results may be incomplete"
                logger.warning(f"Could not load page {page_num} - stopping")
                break

            page_results, next_url = parse_results_page(soup, url)
            self.pages_scraped += 1
            logger.info(f"Page {page_num}: {len(page_results)} organic listings")
            added = 0
            for biz in page_results:
                key = tuple(
                    re.sub(r"\s+", " ", value).strip().casefold()
                    for value in (biz.name, biz.phone, biz.address, biz.city_state)
                )
                if key in seen:
                    self.duplicates_skipped += 1
                else:
                    seen.add(key)
                    all_results.append(biz)
                    added += 1
            self.events.append({"page": page_num, "parsed": len(page_results), "added": added})

            if not next_url or not page_results:
                logger.info("No more pages")
                break

            if next_url in visited:
                self.last_error = "Repeated pagination URL; stopped to prevent a loop"
                break
            if urlsplit(next_url).netloc != urlsplit(self.SEARCH_URL).netloc:
                self.last_error = "Pagination points outside the directory host; stopped"
                break
            url = next_url

            if throttle and page_num < max_pages:
                delay = random.uniform(3, 7)
                logger.info(f"Waiting {delay:.1f}s...")
                time.sleep(delay)

        logger.info(f"Total collected: {len(all_results)}")
        return all_results


# ---------------------------------------------------------------------------
# Demo data generator (no live requests)
# ---------------------------------------------------------------------------
def generate_demo_results(
    category: str, city: str, count: int = 35, *, seed: int = 42
) -> list[Business]:
    """Deterministic fictional businesses with reserved domains and phone numbers."""
    if not 1 <= count <= 1000:
        raise ValueError("Demo count must be between 1 and 1000")
    rng = random.Random(seed)
    city_name = city.split(",")[0].strip()
    locations = {
        "austin": ("512", "78701"),
        "chicago": ("312", "60601"),
        "seattle": ("206", "98101"),
    }
    area, zipcode = locations.get(city_name.lower(), ("202", ""))
    prefixes = [
        "Cedar",
        "Northline",
        "Oak & Co.",
        "Clearwater",
        "Summit",
        "Atlas",
        "Riverside",
        "Heritage",
        "Evergreen",
        "Beacon",
        "Meridian",
        "Parkside",
    ]
    suffixes = {
        "plumbers": ["Plumbing", "Pipe Works", "Plumbing & Drain"],
        "dentists": ["Dental", "Smile Studio", "Family Dentistry"],
        "electricians": ["Electric", "Power Solutions", "Electrical Services"],
    }.get(category.lower(), [category.title(), f"{category.title()} Services"])
    streets = ["Oak Avenue", "Cedar Lane", "Commerce Boulevard", "Market Street", "Riverside Drive"]
    results = []
    for i in range(count):
        name = f"{prefixes[i % len(prefixes)]} {suffixes[(i // len(prefixes)) % len(suffixes)]}"
        slug = re.sub(r"[^a-z0-9]+", "-", name.lower()).strip("-")
        has_rating = i % 7 != 6
        results.append(
            Business(
                name=name,
                phone=f"({area}) 555-{100 + i % 100:04d}",
                address=f"{100 + i * 37} {rng.choice(streets)}",
                city_state=f"{city.strip()} {zipcode}".strip(),
                website=f"https://{slug}.example" if i % 4 != 3 else "",
                category=category.title(),
                rating=f"{rng.choice([3, 3.5, 4, 4.5, 5]):g}/5" if has_rating else "",
                review_count=str(rng.randint(8, 240)) if has_rating else "",
            )
        )
    return results


def run_demo(category="plumbers", city="Austin, TX", count=35, seed=42):
    """Run generated HTML pages through the production pagination and parsing code."""
    from html import escape

    def card(biz, ad=False):
        stars = ""
        if biz.rating:
            value = float(biz.rating.split("/")[0])
            word = ["zero", "one", "two", "three", "four", "five"][int(value)]
            stars = f'<div class="result-rating {word}{" half" if value % 1 else ""}"><span class="count">({biz.review_count})</span></div>'
        website = (
            f'<a class="track-visit-website" href="{escape(biz.website)}?utm_source=demo">Website</a>'
            if biz.website
            else ""
        )
        return (
            f'<div class="result{" flash-endt" if ad else ""}">'
            f'<a class="business-name">{escape(biz.name)}</a>'
            f'<div class="phones">{escape(biz.phone)}</div>'
            f'<div class="street-address">{escape(biz.address)}</div>'
            f'<div class="locality">{escape(biz.city_state)}</div>'
            f'<div class="categories">{escape(biz.category)}</div>{website}{stars}</div>'
        )

    source = generate_demo_results(category, city, count, seed=seed)
    scraper = DirectoryScraper(category, city)
    pages = {}
    total_pages = (count + 11) // 12
    first = scraper._first_url()
    for i in range(total_pages):
        chunk = source[i * 12 : (i + 1) * 12]
        # Intentional repeat at boundaries proves deduplication in the demo.
        if i:
            chunk = [source[i * 12 - 1], *chunk]
        html = "".join(card(b) for b in chunk) + card(source[0], ad=True) * 2
        if i + 1 < total_pages:
            html += f'<a class="next" href="{escape(first)}&amp;page={i + 2}">Next</a>'
        pages[first if i == 0 else f"{first}&page={i + 1}"] = BeautifulSoup(html, "lxml")
    results = scraper.scrape(total_pages, page_loader=pages.get, throttle=False)
    return results, {
        "pages": scraper.pages_scraped,
        "duplicates": scraper.duplicates_skipped,
        "ads": total_pages * 2,
        "events": scraper.events,
        "seed": seed,
        "source": "Synthetic demo • generated HTML • no network requests",
    }


# ---------------------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------------------
HEADERS = [
    "Business Name",
    "Phone",
    "Address",
    "City / State / ZIP",
    "Website",
    "Category",
    "Rating",
    "Review Count",
    "Date Scraped",
]


def export_to_excel(
    businesses: list[Business], output_path, *, source="Directory listings"
) -> None:
    """Export literal text, typed metrics, filters, and frozen headers atomically."""
    if not businesses:
        logger.warning("No results to export")
        return
    rows = []
    for b in businesses:
        values = [
            b.name,
            b.phone,
            b.address,
            b.city_state,
            b.website,
            b.category,
            b.rating,
            b.review_count,
            b.date_scraped,
        ]
        values = [ILLEGAL_CHARACTERS_RE.sub("", value) for value in values]
        with suppress(ValueError):
            values[6] = float(values[6].split("/")[0]) if values[6] else None
        if values[7].isdigit():
            values[7] = int(values[7])
        with suppress(ValueError):
            values[8] = datetime.date.fromisoformat(values[8])
        rows.append(dict(zip(HEADERS, values, strict=True)))
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl", date_format="yyyy-mm-dd") as writer:
        pd.DataFrame(rows).to_excel(writer, index=False, sheet_name="Results")
        wb, ws = writer.book, writer.sheets["Results"]
        wb.properties.title = "Business Directory | Results"
        wb.properties.description = source
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 85
        ws.row_dimensions[1].height = 32
        for cell in ws[1]:
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
            cell.fill = PatternFill("solid", fgColor="123D42")
            cell.alignment = Alignment(vertical="center")
        for row in ws.iter_rows(min_row=2):
            ws.row_dimensions[row[0].row].height = 27
            for cell in row:
                # Untrusted HTML must remain literal text, never an Excel formula.
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11, color="253F43")
                cell.alignment = Alignment(vertical="center")
                if cell.row % 2 == 0:
                    cell.fill = PatternFill("solid", fgColor="EDF6F4")
            row[6].number_format = '0.0" / 5"'
            row[7].number_format = "#,##0"
            row[8].number_format = "yyyy-mm-dd"
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(
                max(len(str(cell.value or "")) for cell in col) + 4, 48
            )
        overview = wb.create_sheet("About this export")
        for row in [
            ["BUSINESS DIRECTORY", "EXPORT NOTES"],
            ["Source", source],
            ["Records", len(businesses)],
            ["Generated", datetime.date.today()],
            ["Missing values", "Blank cells mean the field was not available."],
            ["Ratings", "Numeric values on a five-point scale."],
            ["Demo contacts", "Synthetic exports use fictional numbers and .example domains."],
        ]:
            overview.append(row)
        overview["B4"].number_format = "yyyy-mm-dd"
        overview.column_dimensions["A"].width = 24
        overview.column_dimensions["B"].width = 84
        overview.sheet_view.showGridLines = False
        for row in overview:
            overview.row_dimensions[row[0].row].height = 30
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.font = Font(name="Calibri", size=11, color="253F43")
        for cell in overview[1]:
            cell.fill = PatternFill("solid", fgColor="123D42")
            cell.font = Font(name="Calibri", color="FFFFFF", bold=True, size=12)
    if hasattr(output_path, "write"):
        output_path.write(buffer.getvalue())
    else:
        path = Path(output_path)
        if path.suffix.lower() != ".xlsx":
            raise ValueError("Output filename must end in .xlsx")
        path.parent.mkdir(parents=True, exist_ok=True)
        temp_path = None
        try:
            with tempfile.NamedTemporaryFile(
                dir=path.parent, suffix=".xlsx", delete=False
            ) as handle:
                temp_path = Path(handle.name)
                handle.write(buffer.getvalue())
            os.replace(temp_path, path)
        finally:
            if temp_path and temp_path.exists():
                temp_path.unlink()
    logger.info("Exported %s businesses", len(businesses))


def filename_slug(value: str) -> str:
    """Portable filename component, including Windows reserved-name handling."""
    slug = re.sub(r"[^\w-]+", "_", value, flags=re.UNICODE).strip("_.")[:70] or "results"
    if slug.upper() in {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        *[f"COM{i}" for i in range(1, 10)],
        *[f"LPT{i}" for i in range(1, 10)],
    }:
        slug = f"_{slug}"
    return slug


def positive_int(value):
    number = int(value)
    if number < 1:
        raise argparse.ArgumentTypeError("must be a positive integer")
    return number


def main(argv=None):
    parser = argparse.ArgumentParser(
        description="Business Directory | structured listings to Excel"
    )
    parser.add_argument("--category", default="plumbers")
    parser.add_argument("--city", default="Austin, TX")
    parser.add_argument("--max-pages", type=positive_int, default=5)
    parser.add_argument("--output", help="Output .xlsx path (parent folders are created)")
    parser.add_argument(
        "--demo", action="store_true", help="Parse synthetic HTML without network access"
    )
    parser.add_argument("--count", type=positive_int, default=35, help="Demo records, 1–1000")
    parser.add_argument("--seed", type=int, default=42, help="Reproducible demo seed")
    parser.add_argument(
        "--serve", action="store_true", help="Open the local offline demonstration dashboard"
    )
    parser.add_argument("--port", type=int, default=8765)
    args = parser.parse_args(argv)
    if not args.category.strip() or not args.city.strip():
        parser.error("Category and city must not be blank")
    if args.count > 1000:
        parser.error("Demo count must not exceed 1000")
    if not 1 <= args.port <= 65535:
        parser.error("Port must be between 1 and 65535")
    if args.output and Path(args.output).suffix.lower() != ".xlsx":
        parser.error("Output filename must end in .xlsx")
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(levelname)s] %(message)s")
    try:
        if args.serve:
            from dashboard import serve

            serve(args.port)
            return 0
        output = (
            args.output
            or f"{filename_slug(args.category)}_{filename_slug(args.city.split(',')[0])}{'_demo' if args.demo else ''}.xlsx"
        )
        error = None
        if args.demo:
            results, stats = run_demo(args.category, args.city, args.count, args.seed)
            source = stats["source"]
            logger.info(
                "Demo parsed %s pages; removed %s duplicates and %s ads",
                stats["pages"],
                stats["duplicates"],
                stats["ads"],
            )
        else:
            with DirectoryScraper(args.category, args.city) as scraper:
                results = scraper.scrape(args.max_pages)
                error = scraper.last_error
            source = f"Yellow Pages | {args.category} | {args.city}" + (
                f" | PARTIAL: {error}" if error else ""
            )
        if not results:
            logger.error(error or "No results found")
            return 1
        export_to_excel(results, output, source=source)
        if error:
            print(f"[PARTIAL] {len(results)} businesses saved to {output}. {error}")
            return 2
        print(f"[OK] {len(results)} businesses saved to: {output}")
        return 0
    except (OSError, ValueError) as exc:
        logger.error("%s", exc)
        return 1
    except KeyboardInterrupt:
        logger.info("Stopped")
        return 130


if __name__ == "__main__":
    raise SystemExit(main())
